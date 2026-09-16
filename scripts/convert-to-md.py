#!/usr/bin/env python3
import re
import sys
from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent.parent / "courses"

def is_certificate(path: Path) -> bool:
    return "certificate" in path.stem.lower()


def convert_txt(path: Path) -> str:
    lines = path.read_text(encoding="utf-8").splitlines()
    out = []
    i = 0
    if len(lines) >= 3 and lines[0].strip().startswith("####") and lines[2].strip().startswith("####"):
        title = lines[1].strip().strip("#").strip()
        out.append(f"# {title}")
        i = 3
    while i < len(lines):
        line = lines[i]
        if line.startswith("### "):
            out.append("## " + line[4:])
        elif line.startswith("## "):
            out.append("### " + line[3:])
        else:
            out.append(line)
        i += 1
    text = "\n".join(out).strip() + "\n"
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text


def convert_pptx(path: Path) -> str:
    pres = Presentation(str(path))
    slides = list(pres.slides)
    module_match = re.match(r"Module_(\d+)", path.stem)
    module_no = module_match.group(1) if module_match else "?"
    first_texts = []
    for shape in slides[0].shapes if slides else []:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs).strip()
                if t:
                    first_texts.append(t)
    doc_title = first_texts[0] if first_texts else path.stem.replace("_", " ")
    out = [f"# Module {module_no}: {doc_title}", ""]
    for idx, slide in enumerate(slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        texts.append(t)
        out.append(f"## Slide {idx}")
        if texts:
            out.append(f"**{texts[0]}**")
            out.append("")
            for t in texts[1:]:
                out.append(f"- {t}")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append("")
                out.append("> **Presenter notes:**")
                for nline in notes.splitlines():
                    out.append(f"> {nline}")
        out.append("")
    text = "\n".join(out).strip() + "\n"
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text


def convert_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    title = path.stem.replace("_", " ").replace("-", " ")
    title = re.sub(r"\s+", " ", title).strip()
    out = [f"# {title}", ""]
    for idx, page in enumerate(reader.pages, 1):
        content = (page.extract_text() or "").strip()
        out.append(f"## Page {idx}")
        out.append("")
        out.append(content if content else "*(no extractable text on this page)*")
        out.append("")
    text = "\n".join(out).strip() + "\n"
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text


def main():
    dry_run = "--apply" not in sys.argv
    converted = []
    skipped_certs = []
    for path in sorted(ROOT.rglob("*")):
        if not path.is_file():
            continue
        rel = str(path.relative_to(ROOT.parent))
        if path.suffix.lower() == ".pdf" and is_certificate(path):
            skipped_certs.append(rel)
            continue
        if path.suffix.lower() == ".txt":
            md_text = convert_txt(path)
        elif path.suffix.lower() == ".pptx":
            md_text = convert_pptx(path)
        elif path.suffix.lower() == ".pdf":
            md_text = convert_pdf(path)
        else:
            continue
        md_path = path.with_suffix(".md")
        converted.append((path, md_path, len(md_text)))
        if not dry_run:
            md_path.write_text(md_text, encoding="utf-8")
            path.unlink()

    print(f"{'[DRY RUN] ' if dry_run else ''}Converted: {len(converted)}")
    for src, dst, size in converted:
        print(f"  {src.relative_to(ROOT.parent)} -> {dst.name} ({size} bytes)")
    print(f"\nSkipped certificates (kept as PDF): {len(skipped_certs)}")
    for c in skipped_certs:
        print(f"  {c}")


if __name__ == "__main__":
    main()
